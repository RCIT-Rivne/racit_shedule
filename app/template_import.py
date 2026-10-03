"""Імпорт постійного розкладу з шаблону навчальної частини.

Аркуш «Шаблон» має ту саму сітку, що й щоденні файли «Розклад на ДД_ММ_РРРР (День)»:
рядок 2 — групи, рядок 3 — «Дисципліна | Викладач | Ауд.» трійками з колонки C,
далі по 16 рядків на день (пн … сб), пара — два рядки: чисельник і знаменник.
Порожній рядок знаменника означає «так само, як чисельник», «____» — пари немає.

Аудиторій у шаблоні бракує, тому вони доповнюються з щоденних файлів: з кожного
береться лише день із назви файлу і лише рядок того тижня, що вказаний у C1
(інший рядок щоденного файлу — застарілий залишок).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import openpyxl
from sqlalchemy import select

from .db import Teacher, TimetableEntry, session
from .scheduler.loader import unify_teachers
from .scheduler.models import DAYS, WEEKS, course_of, normalize_spaces, normalize_teacher
from .scheduler.semester import Semester

TEMPLATE_SHEET = "Шаблон"
DAILY_SHEET = "Поточний розклад"
TEACHERS_SHEET = "Викладачі"
FIRST_ROW = 4
ROWS_PER_DAY = 16
PAIRS_PER_DAY = 8
DAILY_NAME = re.compile(r"^Розклад на (\d\d)_(\d\d)_(\d{4}) \((.+)\)\.xlsx$")
SHARED_ROOMS = {"с/з", "спортзал", "тир"}

Key = tuple[int, int, int, str]  # (тиждень, день, пара, група)


@dataclass
class Cell:
    subject: str
    teachers: list[str]
    rooms: list[str]


@dataclass
class Report:
    groups: list[str] = field(default_factory=list)
    cells: dict[Key, Cell] = field(default_factory=dict)
    teachers: dict[str, str | None] = field(default_factory=dict)  # ім'я -> пошта
    rooms_from_daily: int = 0
    daily_files: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def missing_rooms(self) -> list[tuple[Key, Cell]]:
        return sorted(((k, c) for k, c in self.cells.items() if not c.rooms), key=lambda kc: (kc[0][1], kc[0][0], kc[0][2], kc[0][3]))

    def conflicts(self) -> list[str]:
        """Викладач або аудиторія в двох групах одночасно."""
        teachers: dict[tuple, list[str]] = {}
        rooms: dict[tuple, list[str]] = {}
        for (w, d, p, g), c in self.cells.items():
            for t in c.teachers:
                teachers.setdefault((w, d, p, t), []).append(g)
            for r in c.rooms:
                if r.lower() not in SHARED_ROOMS:
                    rooms.setdefault((w, d, p, r), []).append(g)
        out = []
        for kind, index in (("викладач", teachers), ("аудиторія", rooms)):
            for (w, d, p, who), groups in sorted(index.items()):
                if len(groups) > 1:
                    out.append(f"{DAYS[d]}, {WEEKS[w].lower()}, {p} пара: {kind} {who} — {', '.join(groups)}")
        return out

    def lines(self) -> list[str]:
        per_week = [sum(1 for k in self.cells if k[0] == w) for w in (0, 1)]
        with_email = sum(1 for e in self.teachers.values() if e)
        missing = self.missing_rooms()
        out = [
            f"Груп: {len(self.groups)}",
            f"Пар: чисельник {per_week[0]}, знаменник {per_week[1]}",
            f"Викладачів: {len(self.teachers)} (з поштою {with_email}, без пошти {len(self.teachers) - with_email})",
            f"Аудиторій підтягнуто з щоденних файлів: {self.rooms_from_daily} ({len(self.daily_files)} файлів)",
            f"Пар без аудиторії: {len(missing)}",
        ]
        out += [f"  {DAYS[d]}, {WEEKS[w].lower()}, {p} пара, {g}: {c.subject}" for (w, d, p, g), c in missing]
        conflicts = self.conflicts()
        out.append(f"Конфлікти: {len(conflicts)}")
        out += [f"  {c}" for c in conflicts]
        out += [f"Увага: {w}" for w in self.warnings]
        return out


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split(value) -> list[str]:
    return [part.strip() for part in re.split(r"[\n,]", _text(value)) if part.strip()]


def _is_blank(subject: str) -> bool:
    """«____» і «-» — пари немає."""
    return not subject or not subject.strip("_- ")


def _cell(raw: tuple) -> Cell | None:
    subject = normalize_spaces(_text(raw[0]).replace("\n", " "))
    if _is_blank(subject):
        return None
    teachers = [normalize_teacher(t) for t in _split(raw[1]) if not _is_blank(t)]
    return Cell(subject, teachers, [r for r in _split(raw[2]) if not _is_blank(r)])


def _groups(ws) -> dict[int, str]:
    """{колонка «Дисципліна»: група}."""
    return {
        c: normalize_spaces(_text(ws.cell(2, c).value))
        for c in range(3, ws.max_column + 1)
        if ws.cell(2, c).value and _text(ws.cell(3, c).value) == "Дисципліна"
    }


def _rows(ws, day: int, pair: int, col: int) -> tuple[tuple, tuple]:
    top = FIRST_ROW + day * ROWS_PER_DAY + (pair - 1) * 2
    read = lambda r: tuple(ws.cell(r, col + k).value for k in range(3))
    return read(top), read(top + 1)


def read_template(ws) -> tuple[list[str], dict[Key, Cell]]:
    groups = _groups(ws)
    cells: dict[Key, Cell] = {}
    for d in range(len(DAYS)):
        for p in range(1, PAIRS_PER_DAY + 1):
            for col, group in groups.items():
                top_raw, bottom_raw = _rows(ws, d, p, col)
                top = _cell(top_raw)
                if all(_text(v) == "" for v in bottom_raw):
                    bottom = Cell(top.subject, list(top.teachers), list(top.rooms)) if top else None
                else:
                    bottom = _cell(bottom_raw)
                    # Той самий запис у знаменнику без аудиторії — аудиторія як у чисельнику.
                    if bottom and top and not bottom.rooms and (bottom.subject, bottom.teachers) == (top.subject, top.teachers):
                        bottom.rooms = list(top.rooms)
                for week, cell in ((0, top), (1, bottom)):
                    if cell:
                        cells[week, d, p, group] = cell
    return list(groups.values()), cells


def read_daily(path: Path) -> tuple[date, dict[Key, Cell]] | None:
    """Пари дня з назви файлу, лише за типом тижня з C1. None — не будній день."""
    m = DAILY_NAME.match(path.name)
    if not m or m.group(4) not in DAYS:
        return None
    day = DAYS.index(m.group(4))
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[DAILY_SHEET]
        week = 0 if "чисельник" in _text(ws["C1"].value).lower() else 1
        cells: dict[Key, Cell] = {}
        for col, group in _groups(ws).items():
            for p in range(1, PAIRS_PER_DAY + 1):
                top_raw, bottom_raw = _rows(ws, day, p, col)
                raw = bottom_raw if week == 1 and any(_text(v) for v in bottom_raw) else top_raw
                cell = _cell(raw)
                if cell:
                    cells[week, day, p, group] = cell
    finally:
        wb.close()
    return date(int(m.group(3)), int(m.group(2)), int(m.group(1))), cells


def read_teachers_sheet(wb) -> dict[str, str | None]:
    if TEACHERS_SHEET not in wb.sheetnames:
        return {}
    out = {}
    for row in wb[TEACHERS_SHEET].iter_rows(values_only=True):
        name = normalize_teacher(_text(row[0] if row else None))
        email = _text(row[1] if len(row) > 1 else None).lower()
        if name and name.lower() != "викладач":
            out[name] = email if "@" in email else None
    return out


def _unify(cells: list[Cell], known: list[str], warnings: list[str]) -> None:
    """«Сергій ПАНДРАК» -> «Пандрак С. Б.» тією ж логікою, що й у генераторі."""
    unify_teachers([SimpleNamespace(teachers=c.teachers) for c in cells], known, warnings)


def build(path: Path, daily_dir: Path | None = None) -> Report:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        if TEMPLATE_SHEET not in wb.sheetnames:
            raise ValueError(f"У файлі немає аркуша «{TEMPLATE_SHEET}»")
        report = Report()
        report.groups, report.cells = read_template(wb[TEMPLATE_SHEET])
        sheet_teachers = read_teachers_sheet(wb)
    finally:
        wb.close()

    known = list(sheet_teachers)
    _unify(list(report.cells.values()), known, report.warnings)

    daily: dict[Key, Cell] = {}
    if daily_dir is not None:
        files = sorted(
            (r for f in Path(daily_dir).glob("Розклад на *.xlsx") if (r := read_daily(f)) is not None),
            key=lambda r: r[0],
        )
        report.daily_files = [d.isoformat() for d, _ in files]
        for _, cells in files:  # пізніший файл перекриває ранній
            _unify(list(cells.values()), known, [])
            daily.update(cells)

    for key, cell in report.cells.items():
        other = daily.get(key)
        if not cell.rooms and other and other.rooms and (other.subject, other.teachers) == (cell.subject, cell.teachers):
            cell.rooms = list(other.rooms)
            report.rooms_from_daily += 1

    names = {t for c in report.cells.values() for t in c.teachers}
    report.teachers = {n: sheet_teachers.get(n) for n in sorted(names | set(sheet_teachers))}
    missing = sorted(n for n in names if not sheet_teachers.get(n))
    if missing:
        report.warnings.append(f"без пошти на аркуші «{TEACHERS_SHEET}»: {', '.join(missing)}")
    return report


def save_teachers(teachers: dict[str, str | None]) -> None:
    """Додати нових викладачів і оновити пошти (наявні записи не видаляються)."""
    existing = {t.name: t for t in session.scalars(select(Teacher))}
    taken = {t.email for t in existing.values() if t.email}
    for name, email in teachers.items():
        teacher = existing.get(name)
        if teacher is None:
            teacher = existing[name] = Teacher(name=name, aliases=[])
            session.add(teacher)
        if email and teacher.email != email and email not in taken:
            teacher.email = email
            taken.add(email)
    session.flush()


def entries(report: Report) -> list[TimetableEntry]:
    return [
        TimetableEntry(group=g, course=course_of(g), subject=c.subject, teachers=list(c.teachers),
                       rooms=list(c.rooms), week=w, day=d, pair=p)
        for (w, d, p, g), c in sorted(report.cells.items(), key=lambda kc: (report.groups.index(kc[0][3]), kc[0][:3]))
    ]


def import_template(path: Path, semester: Semester, name: str, daily_dir: Path | None = None,
                    dry_run: bool = False) -> Report:
    from .timetable import clear_cache, save_timetable

    report = build(path, daily_dir)
    if not dry_run:
        save_teachers(report.teachers)
        save_timetable(name, f"template:{path.name}", semester, entries(report))
        clear_cache()
    return report
