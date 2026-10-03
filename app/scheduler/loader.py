"""Читання вхідного Excel-файлу (аркуші «Навантаження» та «аудиторії»)."""
from __future__ import annotations

from pathlib import Path

import openpyxl

from .models import (
    Group,
    Lesson,
    ProblemData,
    Room,
    is_full_name,
    normalize_spaces,
    normalize_teacher,
    teacher_key,
)

LOAD_SHEET = "навантаження"
ROOMS_SHEET = "аудиторії"
TEACHERS_SHEET = "викладачі"


class InputError(ValueError):
    pass


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_sheet(wb, name: str):
    for ws in wb.worksheets:
        if ws.title.strip().lower() == name:
            return ws
    return None


def _header_index(header: list[str]) -> dict[str, int]:
    wanted = {
        "group": ("група",),
        "subject": ("дисципліна",),
        "teacher": ("викладач",),
        "per_week": ("кількість пар",),
        "shift": ("зміна",),
    }
    idx = {}
    for key, prefixes in wanted.items():
        for i, h in enumerate(header):
            if any(h.lower().startswith(p) for p in prefixes):
                idx[key] = i
                break
    missing = [k for k in wanted if k not in idx and k != "shift"]
    if missing:
        raise InputError(f"На аркуші «Навантаження» не знайдено колонок: {', '.join(missing)}")
    return idx


def load_lessons(ws, warnings: list[str]) -> tuple[list[Lesson], list[Group]]:
    rows = ws.iter_rows(values_only=True)
    header = [_cell_str(v) for v in next(rows)]
    idx = _header_index(header)

    lessons: list[Lesson] = []
    groups: dict[str, Group] = {}
    for row_no, row in enumerate(rows, start=2):
        group = normalize_spaces(_cell_str(row[idx["group"]]))
        if not group:
            continue
        subject = normalize_spaces(_cell_str(row[idx["subject"]]))
        teachers = [
            normalize_teacher(t) for t in _cell_str(row[idx["teacher"]]).split("\n") if t.strip()
        ]
        try:
            per_week = float(row[idx["per_week"]] or 0)
        except (TypeError, ValueError):
            warnings.append(f"Рядок {row_no}: некоректна кількість пар «{row[idx['per_week']]}», пропущено")
            continue
        shift = 1
        if "shift" in idx and row[idx["shift"]] not in (None, ""):
            shift = int(float(row[idx["shift"]]))
        if shift not in (1, 2):
            warnings.append(f"Рядок {row_no}: зміна {shift} не підтримується, використано 1")
            shift = 1
        if per_week <= 0:
            continue
        if not teachers:
            warnings.append(f"Рядок {row_no}: {group} / {subject} без викладача")

        if group not in groups:
            groups[group] = Group(group, shift)
        elif groups[group].shift != shift:
            warnings.append(
                f"Рядок {row_no}: у групи {group} різні зміни в навантаженні, використано {groups[group].shift}"
            )
            shift = groups[group].shift

        lessons.append(Lesson(len(lessons), group, subject, teachers, per_week, shift))
    return lessons, list(groups.values())


def _room_kind(name: str, note: str) -> tuple[str, bool]:
    n, t = name.lower(), note.lower()
    if "спортзал" in n or n in ("с/з",):
        return "gym", True
    if n == "тир":
        return "range", False
    if n == "н.м":
        return "range_reserve", False
    # Аудиторії НУВГП (напр. «1В») — лише для закріплених правилом room занять.
    if "нувгп" in t:
        return "pinned", False
    if "програміст" in t:
        return "programmers", False
    if "комп" in t:
        return "computer", False
    if "лаборатор" in t:
        return "lab", False
    if "крайн" in t:
        return "reserve", False
    if n in ("бібліотека", "уч.к"):
        return "other", False
    return "regular", False


def load_rooms(ws) -> list[Room]:
    rooms: dict[str, Room] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = _cell_str(row[1] if len(row) > 1 else None)
        if not name:
            continue
        note = _cell_str(row[2] if len(row) > 2 else None)
        if name in rooms:
            continue
        kind, shared = _room_kind(name, note)
        rooms[name] = Room(name, note, kind, shared)
    return list(rooms.values())


def unify_teachers(lessons: list[Lesson], known: list[str], warnings: list[str]) -> None:
    """«Сергій ПАНДРАК» -> «Пандрак С. Б.», якщо такий викладач є в навантаженні
    або на аркуші «Викладачі»; інакше -> «Пандрак С.»."""
    candidates: dict[tuple, set[str]] = {}
    for name in [t for l in lessons for t in l.teachers] + known:
        key = teacher_key(name)
        if key and not is_full_name(name):
            candidates.setdefault(key, set()).add(name)

    renamed: dict[str, str] = {}
    for l in lessons:
        for i, name in enumerate(l.teachers):
            if not is_full_name(name):
                continue
            if name not in renamed:
                first, last = name.split()
                matches = candidates.get(teacher_key(name), set())
                target = next(iter(matches)) if len(matches) == 1 else f"{last.capitalize()} {first[0]}."
                renamed[name] = target
                warnings.append(f"Викладача «{name}» записано як «{target}»")
            l.teachers[i] = renamed[name]


def load_known_teachers(ws) -> list[str]:
    names = []
    for row in ws.iter_rows(values_only=True):
        name = normalize_teacher(_cell_str(row[0] if row else None))
        if name and name.lower() != "викладач":
            names.append(name)
    return names


def load_workbook(path: str | Path) -> ProblemData:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    warnings: list[str] = []
    ws = _find_sheet(wb, LOAD_SHEET) or wb.worksheets[0]
    lessons, groups = load_lessons(ws, warnings)
    if not lessons:
        raise InputError("Навантаження порожнє — немає жодного рядка з групою та дисципліною")

    teachers_ws = _find_sheet(wb, TEACHERS_SHEET)
    unify_teachers(lessons, load_known_teachers(teachers_ws) if teachers_ws is not None else [], warnings)

    rooms_ws = _find_sheet(wb, ROOMS_SHEET)
    rooms = load_rooms(rooms_ws) if rooms_ws is not None else []
    if not rooms:
        warnings.append("Аркуш «аудиторії» не знайдено — аудиторії не призначатимуться")
    wb.close()
    return ProblemData(lessons=lessons, groups=groups, rooms=rooms, warnings=warnings)
