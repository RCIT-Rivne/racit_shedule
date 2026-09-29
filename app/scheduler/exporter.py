"""Експорт розкладу в Excel у форматі, схожому на шаблон коледжу.

Кожна пара займає два рядки: верхній — чисельник, нижній — знаменник.
Якщо заняття однакове в обох тижнях, рядки об'єднуються.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .metrics import quality
from .models import DAYS, SATURDAY_NAME, TOTAL_PAIRS, WEEKS, Group, Schedule

FONT = "Times New Roman"
THIN = Side(style="thin")
THICK = Side(style="medium")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
DAY_FILL = PatternFill("solid", fgColor="F2F2F2")
ERROR_FILL = PatternFill("solid", fgColor="F8CBAD")
WARN_FILL = PatternFill("solid", fgColor="FFE699")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

ROWS_PER_DAY = TOTAL_PAIRS * 2
FIRST_ROW = 4


def _style(cell, bold=False, size=10, align=CENTER, fill=None):
    cell.font = Font(name=FONT, size=size, bold=bold)
    cell.alignment = align
    if fill:
        cell.fill = fill


def _box(ws, min_row, min_col, max_row, max_col, bottom=THIN):
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(r, c).border = Border(
                left=THICK if c == min_col else THIN,
                right=THIN,
                top=THIN,
                bottom=bottom if r == max_row else THIN,
            )


def _cell_values(schedule: Schedule, placement) -> tuple[str, str, str]:
    lesson = schedule.lesson(placement.lesson_id)
    return lesson.subject, "\n".join(lesson.teachers), "\n".join(placement.rooms)


def _print_setup(ws, title_rows: str):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = title_rows
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4


def _write_group_sheet(ws, schedule: Schedule, groups: list[Group], title: str, hide_unused: bool = False):
    ws.sheet_view.zoomScale = 70
    ws.freeze_panes = ws.cell(FIRST_ROW, 3)
    ws.cell(1, 1, f"{title} (верхній рядок пари — чисельник, нижній — знаменник)")
    _style(ws.cell(1, 1), bold=True, size=14, align=Alignment(vertical="center"))
    ws.row_dimensions[1].height = 24
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2)
    ws.cell(2, 1, "Група")
    ws.cell(3, 1, "День")
    ws.cell(3, 2, "Пара")
    for c in (ws.cell(2, 1), ws.cell(3, 1), ws.cell(3, 2)):
        _style(c, bold=True, size=11, fill=HEADER_FILL)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 6
    for gi, group in enumerate(groups):
        col = 3 + gi * 3
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 2)
        ws.cell(2, col, group.name)
        _style(ws.cell(2, col), bold=True, size=12, fill=HEADER_FILL)
        for i, name in enumerate(("Дисципліна", "Викладач", "Ауд.")):
            ws.cell(3, col + i, name)
            _style(ws.cell(3, col + i), bold=True, fill=HEADER_FILL)
        ws.column_dimensions[get_column_letter(col)].width = 26
        ws.column_dimensions[get_column_letter(col + 1)].width = 18
        ws.column_dimensions[get_column_letter(col + 2)].width = 6
    _box(ws, 2, 1, 3, 2 + 3 * len(groups), bottom=THICK)
    _print_setup(ws, "1:3")
    names = {g.name for g in groups}
    used_pairs = {pair for g in groups for pair in g.pair_numbers()} | {
        p.pair for p in schedule.placements if schedule.lesson(p.lesson_id).group in names
    }

    for d, day in enumerate(schedule.data.days):
        top = FIRST_ROW + d * ROWS_PER_DAY
        ws.merge_cells(start_row=top, start_column=1, end_row=top + ROWS_PER_DAY - 1, end_column=1)
        ws.cell(top, 1, day)
        _style(ws.cell(top, 1), bold=True, size=12, align=Alignment(
            horizontal="center", vertical="center", text_rotation=90), fill=DAY_FILL)
        for pair in range(1, TOTAL_PAIRS + 1):
            r = top + (pair - 1) * 2
            ws.merge_cells(start_row=r, start_column=2, end_row=r + 1, end_column=2)
            ws.cell(r, 2, pair)
            _style(ws.cell(r, 2), bold=True, size=11, fill=DAY_FILL)
            ws.row_dimensions[r].height = 27
            ws.row_dimensions[r + 1].height = 27
            if hide_unused and pair not in used_pairs:
                ws.row_dimensions[r].hidden = True
                ws.row_dimensions[r + 1].hidden = True

            for gi, group in enumerate(groups):
                col = 3 + gi * 3
                num = schedule.cell(group.name, 0, d, pair)
                den = schedule.cell(group.name, 1, d, pair)
                vnum = _cell_values(schedule, num) if num else None
                vden = _cell_values(schedule, den) if den else None
                if vnum and vnum == vden:
                    for i, v in enumerate(vnum):
                        ws.merge_cells(start_row=r, start_column=col + i, end_row=r + 1, end_column=col + i)
                        ws.cell(r, col + i, v)
                else:
                    for row, vals in ((r, vnum), (r + 1, vden)):
                        if vals is None and (vnum or vden):
                            vals = ("—", "", "")
                        for i, v in enumerate(vals or ()):
                            ws.cell(row, col + i, v)
                for rr in (r, r + 1):
                    for i in range(3):
                        _style(ws.cell(rr, col + i), size=10, align=LEFT if i < 2 else CENTER)
        _box(ws, top, 1, top + ROWS_PER_DAY - 1, 2 + 3 * len(groups), bottom=THICK)


def _write_teachers(ws, schedule: Schedule):
    """Зведення по викладачах: рядок на тиждень, стовпчик на пару."""
    days = schedule.data.days
    ws.sheet_view.zoomScale = 80
    ws.freeze_panes = "C3"
    ws.cell(1, 1, "Викладач")
    ws.cell(1, 2, "Тиждень")
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    for d, day in enumerate(days):
        c = 3 + d * TOTAL_PAIRS
        ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + TOTAL_PAIRS - 1)
        ws.cell(1, c, day)
        for p in range(TOTAL_PAIRS):
            ws.cell(2, c + p, p + 1)
            ws.column_dimensions[get_column_letter(c + p)].width = 9
    for row in ws.iter_rows(min_row=1, max_row=2, max_col=2 + len(days) * TOTAL_PAIRS):
        for cell in row:
            _style(cell, bold=True, fill=HEADER_FILL)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 11
    _print_setup(ws, "1:2")

    grid: dict[tuple, str] = {}
    for p in schedule.placements:
        l = schedule.lesson(p.lesson_id)
        for i, t in enumerate(l.teachers):
            # У підгрупах кожен викладач — у своїй аудиторії.
            room = p.rooms[i] if i < len(p.rooms) else ", ".join(p.rooms)
            grid[t, p.week, p.day, p.pair] = f"{l.group} ({room})" if room else l.group

    row = 3
    for teacher in schedule.data.teachers:
        ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=1)
        ws.cell(row, 1, teacher)
        _style(ws.cell(row, 1), bold=True, align=LEFT)
        for w in range(len(WEEKS)):
            ws.cell(row + w, 2, WEEKS[w])
            _style(ws.cell(row + w, 2), size=9)
            for d in range(len(days)):
                for p in range(TOTAL_PAIRS):
                    cell = ws.cell(row + w, 3 + d * TOTAL_PAIRS + p, grid.get((teacher, w, d, p + 1)))
                    _style(cell, size=9)
        _box(ws, row, 1, row + 1, 2 + len(days) * TOTAL_PAIRS)
        row += 2


def _write_load(ws, schedule: Schedule):
    """Контроль навчального плану: план vs поставлено по тижнях."""
    headers = ["Група", "Дисципліна", "Викладач", "План, пар/тиж", "Чисельник", "Знаменник"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
        _style(ws.cell(1, i), bold=True, fill=HEADER_FILL)
    counts = defaultdict(int)
    for p in schedule.placements:
        counts[p.lesson_id, p.week] += 1
    for r, l in enumerate(schedule.data.lessons, 2):
        values = [l.group, l.subject, ", ".join(l.teachers), l.per_week, counts[l.id, 0], counts[l.id, 1]]
        for i, v in enumerate(values, 1):
            ws.cell(r, i, v)
            _style(ws.cell(r, i), align=LEFT if i <= 3 else CENTER)
    for col, width in zip("ABCDEF", (12, 50, 36, 14, 12, 12)):
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = f"A1:F{len(schedule.data.lessons) + 1}"
    ws.freeze_panes = "A2"


def _write_issues(ws, schedule: Schedule):
    headers = ["Рівень", "Правило", "Опис"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
        _style(ws.cell(1, i), bold=True, fill=HEADER_FILL)
    ws.cell(2, 1, "Статус розв'язувача")
    ws.cell(2, 2, schedule.status)
    ws.cell(2, 3, f"штраф {schedule.objective}, час {schedule.solve_seconds} с")
    labels = {"error": "Помилка", "warning": "Попередження", "info": "Інфо"}
    fills = {"error": ERROR_FILL, "warning": WARN_FILL}
    for r, issue in enumerate(schedule.issues, 3):
        for i, v in enumerate((labels[issue.severity], issue.rule, issue.message), 1):
            ws.cell(r, i, v)
            _style(ws.cell(r, i), align=LEFT, fill=fills.get(issue.severity))
    if not schedule.issues:
        ws.cell(3, 1, "Порушень правил не знайдено")

    row = max(ws.max_row, 3) + 2
    for title, values in (("Якість розкладу", quality(schedule)), ("Штраф за правилами", schedule.penalty_breakdown)):
        ws.cell(row, 1, title)
        _style(ws.cell(row, 1), bold=True, align=LEFT, fill=HEADER_FILL)
        row += 1
        for k, v in values.items():
            ws.cell(row, 2, k)
            ws.cell(row, 3, str(v))
            _style(ws.cell(row, 2), align=LEFT)
            _style(ws.cell(row, 3), align=LEFT)
            row += 1
        row += 1
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 100


def _write_saturdays(ws, schedule: Schedule):
    """Календар навчальних субот: за розкладом якого дня і тижня."""
    headers = ["Дата", "Тиждень", "Навчання за розкладом"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
        _style(ws.cell(1, i), bold=True, fill=HEADER_FILL)
    for r, s in enumerate(schedule.semester.saturdays_list(), 2):
        values = (s["date"].strftime("%d.%m.%Y"), WEEKS[s["week"]], DAYS[s["day"]])
        for i, v in enumerate(values, 1):
            ws.cell(r, i, v)
            _style(ws.cell(r, i), align=LEFT)
    for col, width in zip("ABC", (14, 14, 26)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


def build_workbook(schedule: Schedule) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Розклад"
    groups = schedule.data.groups
    _write_group_sheet(ws, schedule, groups, "Розклад занять")

    by_course: dict[int, list[Group]] = defaultdict(list)
    for g in groups:
        by_course[g.course].append(g)
    for course in sorted(by_course):
        title = f"{course} курс" if course else "Інші групи"
        _write_group_sheet(
            wb.create_sheet(title), schedule, by_course[course], f"Розклад занять — {title}", hide_unused=True
        )

    _write_teachers(wb.create_sheet("Викладачі"), schedule)
    _write_load(wb.create_sheet("Навантаження"), schedule)
    if schedule.semester and schedule.semester.saturdays:
        _write_saturdays(wb.create_sheet(f"{SATURDAY_NAME}и"), schedule)
    _write_issues(wb.create_sheet("Перевірка"), schedule)
    return wb


def export_xlsx(schedule: Schedule, target: str | Path | BytesIO) -> None:
    build_workbook(schedule).save(target)


def day_title(schedule: Schedule, day: date) -> str:
    weekday = DAYS[day.weekday()] if day.weekday() < 5 else SATURDAY_NAME
    return f"Розклад на {day.strftime('%d.%m.%Y')} ({weekday})"


def build_day_workbook(schedule: Schedule, day: date) -> Workbook:
    """Розклад на конкретну дату: один день, лише актуальний тиждень, усі групи."""
    resolved = schedule.semester.resolve(day)
    wb = Workbook()
    ws = wb.active
    ws.title = day.strftime("%d.%m.%Y")
    ws.cell(1, 1, f"{day_title(schedule, day)} — {schedule.semester.describe(day)}")
    _style(ws.cell(1, 1), bold=True, size=14, align=Alignment(vertical="center"))
    ws.row_dimensions[1].height = 24
    if resolved is None:
        return wb
    d, week = resolved
    groups = schedule.data.groups
    ws.cell(2, 1, "Пара")
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)
    _style(ws.cell(2, 1), bold=True, fill=HEADER_FILL)
    ws.column_dimensions["A"].width = 6
    for gi, group in enumerate(groups):
        col = 2 + gi * 3
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 2)
        ws.cell(2, col, group.name)
        _style(ws.cell(2, col), bold=True, size=12, fill=HEADER_FILL)
        for i, name in enumerate(("Дисципліна", "Викладач", "Ауд.")):
            ws.cell(3, col + i, name)
            _style(ws.cell(3, col + i), bold=True, fill=HEADER_FILL)
        ws.column_dimensions[get_column_letter(col)].width = 26
        ws.column_dimensions[get_column_letter(col + 1)].width = 18
        ws.column_dimensions[get_column_letter(col + 2)].width = 6
    for pair in range(1, TOTAL_PAIRS + 1):
        r = 3 + pair
        ws.cell(r, 1, pair)
        _style(ws.cell(r, 1), bold=True, size=11, fill=DAY_FILL)
        ws.row_dimensions[r].height = 36
        for gi, group in enumerate(groups):
            col = 2 + gi * 3
            placement = schedule.cell(group.name, week, d, pair)
            values = _cell_values(schedule, placement) if placement else ("", "", "")
            for i, v in enumerate(values):
                ws.cell(r, col + i, v)
                _style(ws.cell(r, col + i), align=LEFT if i < 2 else CENTER)
    _box(ws, 2, 1, 3 + TOTAL_PAIRS, 1 + 3 * len(groups))
    ws.freeze_panes = "B4"
    _print_setup(ws, "1:3")
    return wb
