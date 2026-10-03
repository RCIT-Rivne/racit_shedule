"""Перевірка файлу навантаження до генерації.

Генератор на некоректному файлі або мовчки пропускає рядки, або шукає розклад
хвилинами й не знаходить. Тут ті самі дані читаються заздалегідь і кожна
проблема описується з номером рядка — помилки зупиняють генерацію.
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .scheduler.loader import (
    LOAD_SHEET,
    ROOMS_SHEET,
    InputError,
    _cell_str,
    _find_sheet,
    _header_index,
    load_workbook,
)
from .scheduler.models import DAYS, normalize_spaces
from .scheduler.solver import SolverConfig
from .scheduler.validator import precheck


@dataclass
class InputCheck:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    groups: int = 0
    teachers: int = 0
    lessons: int = 0
    rooms: int = 0

    @property
    def ok(self) -> bool:
        return not self.errors


def _rows(path: Path, result: InputCheck) -> None:
    """Порядкова перевірка аркуша «Навантаження»."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = _find_sheet(wb, LOAD_SHEET)
        if ws is None:
            result.errors.append("Немає аркуша «Навантаження»")
            return
        if _find_sheet(wb, ROOMS_SHEET) is None:
            result.warnings.append("Немає аркуша «аудиторії» — аудиторії не призначатимуться")
        rows = ws.iter_rows(values_only=True)
        header = [_cell_str(v) for v in next(rows, ())]
        try:
            idx = _header_index(header)
        except InputError as e:
            result.errors.append(str(e))
            return
        shifts: dict[str, tuple[int, int]] = {}
        for row_no, row in enumerate(rows, start=2):
            cell = lambda key: row[idx[key]] if key in idx and idx[key] < len(row) else None
            group = normalize_spaces(_cell_str(cell("group")))
            subject = normalize_spaces(_cell_str(cell("subject")))
            if not group and not subject:
                continue
            where = f"Рядок {row_no}"
            if not group or not subject:
                result.errors.append(f"{where}: не вказано {'групу' if not group else 'дисципліну'}")
                continue
            if not _cell_str(cell("teacher")):
                result.errors.append(f"{where}: {group} / {subject} — не вказано викладача")
            raw = cell("per_week")
            try:
                per_week = float(raw)
            except (TypeError, ValueError):
                result.errors.append(f"{where}: {group} / {subject} — кількість пар «{_cell_str(raw)}» не число")
                continue
            if per_week < 0 or (per_week * 2) % 1:
                result.errors.append(f"{where}: {group} / {subject} — кількість пар {per_week:g}: "
                                     "потрібне ціле число або з половиною (1,5)")
            raw_shift = _cell_str(cell("shift")) or "1"
            if raw_shift not in ("1", "2"):
                result.errors.append(f"{where}: {group} — зміна «{raw_shift}», потрібно 1 або 2")
                continue
            first = shifts.setdefault(group, (int(raw_shift), row_no))
            if first[0] != int(raw_shift):
                result.errors.append(f"{where}: {group} — зміна {raw_shift}, а в рядку {first[1]} — {first[0]}")
    finally:
        wb.close()


def check_input(path: str | Path) -> InputCheck:
    path = Path(path)
    result = InputCheck()
    try:
        _rows(path, result)
    except Exception as e:  # пошкоджений або не xlsx-файл
        result.errors.append(f"Файл не вдалося прочитати: {e}")
        return result
    if result.errors:
        return result

    data = load_workbook(path)
    result.warnings += data.warnings
    result.groups, result.lessons, result.rooms = len(data.groups), len(data.lessons), len(data.rooms)
    result.teachers = len({t for l in data.lessons for t in l.teachers})
    config = SolverConfig()
    result.errors += [
        f"{i.rule}: {i.message}"
        for i in precheck(data, len(DAYS), config.max_pairs_per_day, config.max_teacher_pairs_per_day)
    ]

    # Дубль рядка (група + дисципліна + викладач) — генератор поставить пари двічі.
    seen: dict[tuple, int] = defaultdict(int)
    for l in data.lessons:
        seen[l.group, l.subject, tuple(l.teachers)] += 1
    result.warnings += [f"{g} / {s} ({', '.join(t)}) записано {n} рази" for (g, s, t), n in seen.items() if n > 1]
    return result
