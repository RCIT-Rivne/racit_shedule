"""Імпорт постійного розкладу з шаблону навчальної частини (аркуш «Шаблон»)."""
from datetime import date

import openpyxl
import pytest

from app.scheduler.semester import Semester
from app.template_import import FIRST_ROW, ROWS_PER_DAY, build

SEMESTER = Semester(date(2026, 9, 1), date(2026, 12, 31))
GROUPS = ["КН-1/1", "ІПЗ-2/1"]
MON, TUE = 0, 1


def _sheet(wb, title, cells, week_label="чисельнику"):
    """cells: {(день, пара, 0|1 — верхній/нижній рядок, група): (дисципліна, викладач, ауд.)}"""
    ws = wb.create_sheet(title)
    ws["C1"] = f"Навчання здійснюється по {week_label}"
    for i, g in enumerate(GROUPS):
        col = 3 + i * 3
        ws.cell(2, col, g)
        for k, h in enumerate(("Дисципліна", "Викладач", "Ауд.")):
            ws.cell(3, col + k, h)
    for (d, p, half, g), values in cells.items():
        row = FIRST_ROW + d * ROWS_PER_DAY + (p - 1) * 2 + half
        col = 3 + GROUPS.index(g) * 3
        for k, v in enumerate(values):
            ws.cell(row, col + k, v)
    return ws


def _template(tmp_path, cells, teachers=()):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Шаблон", cells)
    ws = wb.create_sheet("Викладачі")
    for name, email in teachers:
        ws.append([name, email])
    path = tmp_path / "Розклад (Шаблон).xlsx"
    wb.save(path)
    return path


def _daily(tmp_path, name, day, cells, week_label):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Поточний розклад", {(day, p, h, g): v for (p, h, g), v in cells.items()}, week_label)
    wb.save(tmp_path / name)


def test_template_rows_and_normalisation(tmp_path):
    path = _template(tmp_path, {
        # однаково в обох тижнях: нижній рядок порожній, аудиторія з Excel — число
        (MON, 1, 0, "КН-1/1"): ("Математика", "Бойко В.В.", 26.0),
        # різне: знаменник — «пари немає»
        (MON, 2, 0, "КН-1/1"): ("Фізика", "Криволисова К. А.", None),
        (MON, 2, 1, "КН-1/1"): ("________________", "________________", None),
        # підгрупи + «Ім'я ПРІЗВИЩЕ»
        (TUE, 3, 0, "ІПЗ-2/1"): ("Інформатика", "Назаров А. Л.\nСергій ПАНДРАК", "26\n27"),
        # лише знаменник
        (TUE, 8, 1, "ІПЗ-2/1"): ("Захист України", "Ойцюсь А. М.", "тир"),
    }, teachers=[("Пандрак С. Б.", "spandrak@rcit.ukr.education"), ("Бойко В. В.", "VBoiko@rcit.ukr.education")])

    report = build(path)
    c = report.cells
    assert report.groups == GROUPS
    assert c[0, MON, 1, "КН-1/1"].rooms == ["26"] and c[1, MON, 1, "КН-1/1"].rooms == ["26"]
    assert c[0, MON, 1, "КН-1/1"].teachers == ["Бойко В. В."]
    assert (1, MON, 2, "КН-1/1") not in c and c[0, MON, 2, "КН-1/1"].subject == "Фізика"
    sub = c[0, TUE, 3, "ІПЗ-2/1"]
    assert sub.teachers == ["Назаров А. Л.", "Пандрак С. Б."] and sub.rooms == ["26", "27"]
    assert (0, TUE, 8, "ІПЗ-2/1") not in c and c[1, TUE, 8, "ІПЗ-2/1"].rooms == ["тир"]
    assert report.teachers["Бойко В. В."] == "vboiko@rcit.ukr.education"
    assert report.teachers["Назаров А. Л."] is None
    assert c[0, MON, 2, "КН-1/1"].rooms == []
    assert len(report.missing_rooms()) == 1


def test_rooms_from_daily_only_active_week(tmp_path):
    path = _template(tmp_path, {
        (MON, 1, 0, "КН-1/1"): ("Математика", "Бойко В. В.", None),
        (MON, 1, 1, "КН-1/1"): ("Хімія", "Федорчук Р. Ю.", None),
    })
    # Тиждень-чисельник: верхній рядок — актуальний; нижній — застарілий залишок.
    _daily(tmp_path, "Розклад на 28_09_2026 (Понеділок).xlsx", MON, {
        (1, 0, "КН-1/1"): ("Математика", "Бойко В. В.", 19),
        (1, 1, "КН-1/1"): ("Хімія", "Федорчук Р. Ю.", 99),
    }, "чисельнику")
    report = build(path, tmp_path)
    assert report.cells[0, MON, 1, "КН-1/1"].rooms == ["19"]
    assert report.cells[1, MON, 1, "КН-1/1"].rooms == []

    # Тиждень-знаменник: актуальний нижній рядок.
    _daily(tmp_path, "Розклад на 21_09_2026 (Понеділок).xlsx", MON, {
        (1, 0, "КН-1/1"): ("Математика", "Бойко В. В.", 77),
        (1, 1, "КН-1/1"): ("Хімія", "Федорчук Р. Ю.", 36),
    }, "знаменнику")
    report = build(path, tmp_path)
    assert report.cells[0, MON, 1, "КН-1/1"].rooms == ["19"]
    assert report.cells[1, MON, 1, "КН-1/1"].rooms == ["36"]


def test_daily_room_ignored_when_pair_differs(tmp_path):
    path = _template(tmp_path, {(MON, 1, 0, "КН-1/1"): ("Математика", "Бойко В. В.", None)})
    _daily(tmp_path, "Розклад на 28_09_2026 (Понеділок).xlsx", MON,
           {(1, 0, "КН-1/1"): ("Фізика", "Криволисова К. А.", 25)}, "чисельнику")
    assert build(path, tmp_path).cells[0, MON, 1, "КН-1/1"].rooms == []


def test_conflicts_reported(tmp_path):
    path = _template(tmp_path, {
        (MON, 3, 0, "КН-1/1"): ("ООП", "Ніколаєнко А. І.", 9),
        (MON, 3, 0, "ІПЗ-2/1"): ("Бази даних", "Ніколаєнко А. І.", 9),
    })
    conflicts = build(path).conflicts()
    assert any("викладач Ніколаєнко" in c for c in conflicts)
    assert any("аудиторія 9" in c for c in conflicts)


@pytest.fixture()
def app(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    from app import create_app, db

    app = create_app()
    yield app
    db.session.remove()


def test_import_saves_active_timetable_and_teachers(app, tmp_path):
    from app import db
    from app.template_import import import_template
    from app.timetable import day_slots, for_teacher

    path = _template(tmp_path, {(MON, 1, 0, "КН-1/1"): ("Математика", "Бойко В. В.", 26)},
                     teachers=[("Бойко В. В.", "vboiko@rcit.ukr.education")])
    with app.app_context():
        dry = import_template(path, SEMESTER, "Пробний", dry_run=True)
        assert dry.cells and db.active_timetable() is None

        import_template(path, SEMESTER, "Шаблон")
        tt = db.active_timetable()
        assert tt.name == "Шаблон" and tt.source_job == "template:Розклад (Шаблон).xlsx"
        assert len(tt.entries) == 2  # чисельник і знаменник
        assert db.teacher_by_email("VBoiko@rcit.ukr.education").name == "Бойко В. В."
        slots = for_teacher(day_slots(tt, date(2026, 10, 5)), "Бойко В. В.")
        assert [(s.pair, s.group, s.rooms, s.time) for s in slots] == [(1, "КН-1/1", ["26"], ("08:30", "09:50"))]

        import_template(path, SEMESTER, "Вдруге")
        assert db.active_timetable().name == "Вдруге"
        assert db.session.query(db.Teacher).count() == 1


def test_generator_hidden_by_default(app):
    c = app.test_client()
    assert c.get("/admin/generate").status_code == 404
