"""Новий семестр: шаблони файлів, перевірка навантаження, імпорт шаблону з адмінки."""
from io import BytesIO

import openpyxl
import pytest

from app.input_check import check_input
from app.input_templates import grid_template, load_template
from app.template_import import build


def _save(buf: BytesIO, path):
    path.write_bytes(buf.getvalue())
    return path


def _load_file(tmp_path, rows, rooms=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Навантаження"
    ws.append(["Група", "Дисципліна", "Викладач", "Кількість пар на тиждень", "Зміна"])
    for r in rows:
        ws.append(list(r))
    if rooms:
        wb.create_sheet("аудиторії").append(["№", "Аудиторія", "Примітка"])
    path = tmp_path / "load.xlsx"
    wb.save(path)
    return path


def test_example_load_template_passes_check_and_generates(tmp_path):
    from app.scheduler import SolverConfig, generate

    path = _save(load_template(), tmp_path / "example.xlsx")
    check = check_input(path)
    assert check.ok, check.errors
    assert (check.groups, check.rooms) == (3, 9)
    schedule = generate(path, SolverConfig(time_limit=30))
    assert schedule.placements
    assert not [i for i in schedule.issues if i.severity == "error"]


def test_empty_load_template_has_headers_only(tmp_path):
    wb = openpyxl.load_workbook(_save(load_template(example=False), tmp_path / "empty.xlsx"))
    assert wb.sheetnames == ["Інструкція", "Навантаження", "аудиторії", "Викладачі"]
    assert wb["Навантаження"].max_row == 1


def test_check_reports_row_errors(tmp_path):
    path = _load_file(tmp_path, [
        ("КН-1/1", "Математика", "Бойко В. В.", 2, 1),
        ("КН-1/1", "Фізика", "", 2, 1),             # 3: без викладача
        ("КН-1/1", "Хімія", "Федорчук Р. Ю.", "два", 1),  # 4: не число
        ("КН-1/1", "Біологія", "Багнюк А. В.", 1.3, 1),   # 5: не половинка
        ("КН-1/1", "Історія", "Оніщук К. І.", 1, 3),      # 6: зміна 3
        ("КН-1/1", "Право", "Оніщук К. І.", 1, 2),        # 7: інша зміна групи
        ("", "Географія", "Оніщук К. І.", 1, 1),          # 8: без групи
    ])
    errors = check_input(path).errors
    for row in (3, 4, 5, 6, 7, 8):
        assert any(e.startswith(f"Рядок {row}:") for e in errors), (row, errors)


def test_check_reports_overload(tmp_path):
    path = _load_file(tmp_path, [("ГРС-3/2", f"Дисципліна {i}", f"Викладач {i}. А.", 2, 1) for i in range(13)])
    check = check_input(path)
    assert any("ГРС-3/2: 26 пар" in e for e in check.errors)


def test_check_missing_sheet(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Інше"
    wb.save(tmp_path / "x.xlsx")
    assert check_input(tmp_path / "x.xlsx").errors == ["Немає аркуша «Навантаження»"]


def test_grid_template_roundtrip(tmp_path):
    from types import SimpleNamespace as E

    entries = [
        E(week=0, day=0, pair=1, group="КН-1/1", subject="Математика", teachers=["Бойко В. В."], rooms=["19"]),
        E(week=1, day=0, pair=1, group="КН-1/1", subject="Математика", teachers=["Бойко В. В."], rooms=["19"]),
        E(week=0, day=1, pair=2, group="ІПЗ-2/1", subject="Фізика", teachers=["Криволисова К. А."], rooms=[]),
        E(week=1, day=2, pair=5, group="ІПЗ-2/1", subject="Інформатика", teachers=["Назаров А. Л.", "Сальчук А. В."],
          rooms=["26", "27"]),
    ]
    path = _save(grid_template(["КН-1/1", "ІПЗ-2/1"], entries, [("Бойко В. В.", "b@x.ua")]), tmp_path / "grid.xlsx")
    report = build(path)
    got = {k: (c.subject, c.teachers, c.rooms) for k, c in report.cells.items()}
    assert got == {(e.week, e.day, e.pair, e.group): (e.subject, e.teachers, e.rooms) for e in entries}
    assert report.teachers["Бойко В. В."] == "b@x.ua"


@pytest.fixture()
def app(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("ADMIN_EMAIL", "admin@test.ua")
    monkeypatch.setenv("ADMIN_PASSWORD", "adminpass1")
    from app import create_app, db

    app = create_app()
    yield app
    db.session.remove()


@pytest.fixture()
def admin(app):
    c = app.test_client()
    with c.session_transaction() as s:
        s["csrf"] = "t"
    c.post("/login", data={"email": "admin@test.ua", "password": "adminpass1", "csrf": "t"})
    with c.session_transaction() as s:
        s["csrf"] = "t"
    return c


def test_semester_page_and_downloads(admin):
    assert admin.get("/admin/semester").status_code == 200
    for url in ("/admin/semester/load-template.xlsx", "/admin/semester/load-template.xlsx?empty=1",
                "/admin/semester/grid-template.xlsx", "/admin/semester/grid-template.xlsx?filled=1"):
        r = admin.get(url)
        assert r.status_code == 200 and r.data[:2] == b"PK", url


def test_import_template_from_admin(admin, app, tmp_path):
    from types import SimpleNamespace as E

    from app import db

    grid = grid_template(["КН-1/1"], [E(week=0, day=0, pair=1, group="КН-1/1", subject="Математика",
                                        teachers=["Бойко В. В."], rooms=["19"])])
    r = admin.post("/admin/semester/preview", data={
        "csrf": "t", "start": "2027-01-12", "end": "2027-06-30", "first_week": "0", "saturdays": "on",
        "template": (grid, "Розклад (Шаблон).xlsx"),
    }, content_type="multipart/form-data")
    assert r.status_code == 200 and "Перевірка шаблону" in r.text
    with app.app_context():
        assert db.active_timetable() is None
    token = r.text.split("/admin/semester/import/")[1].split('"')[0]
    r = admin.post(f"/admin/semester/import/{token}", data={"csrf": "t", "name": "ІІ семестр"})
    assert r.status_code == 302
    with app.app_context():
        tt = db.active_timetable()
        assert tt.name == "ІІ семестр" and str(tt.semester_start) == "2027-01-12" and len(tt.entries) == 1  # лише чисельник
    assert admin.post(f"/admin/semester/import/{token}", data={"csrf": "t"}).status_code == 404


def test_generate_refuses_invalid_file(admin, tmp_path):
    path = _load_file(tmp_path, [("КН-1/1", "Хімія", "Федорчук Р. Ю.", "два", 1)])
    r = admin.post("/admin/generate", data={"csrf": "t", "action": "generate", "file": (path.open("rb"), "load.xlsx")},
                   content_type="multipart/form-data")
    assert r.status_code == 400 and "не число" in r.text and "генерацію не запущено" in r.text


def test_generate_check_only(admin, tmp_path):
    path = _save(load_template(), tmp_path / "example.xlsx")
    r = admin.post("/admin/generate", data={"csrf": "t", "action": "check", "file": (path.open("rb"), "example.xlsx")},
                   content_type="multipart/form-data")
    assert r.status_code == 200 and "Помилок немає" in r.text
