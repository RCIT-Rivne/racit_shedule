from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from app.scheduler import Semester, SolverConfig, generate
from app.scheduler.exporter import build_day_workbook, build_workbook
from app.scheduler.loader import unify_teachers
from app.scheduler.models import Group, Lesson, ProblemData, Room, normalize_teacher
from app.scheduler.rooms import assign_rooms
from app.scheduler.solver import solve
from app.scheduler.validator import validate

SAMPLE = Path(__file__).resolve().parent.parent / "info" / "Розклад на 2026 р (інформація).xlsx"
FAST = SolverConfig(time_limit=10)


def small_problem() -> ProblemData:
    """Дві групи, спільні викладачі, дробове навантаження, підгрупи."""
    lessons = []
    plan = [
        ("A-1/1", "Математика", ["Іванов І. І."], 2.5),
        ("A-1/1", "Фізика", ["Петров П. П."], 2),
        ("A-1/1", "Історія", ["Сидоренко С. С."], 1.5),
        ("A-1/1", "Іноземна мова", ["Коваль К. К.", "Мороз М. М."], 2),
        ("A-1/1", "Фізична культура", ["Бондар Б. Б."], 2),
        ("A-1/1", "Інформатика", ["Шевченко Т. Г."], 2),
        ("A-1/1", "Хімія", ["Петров П. П."], 2),
        ("A-1/1", "Біологія", ["Сидоренко С. С."], 2),
        ("B-1/1", "Математика", ["Іванов І. І."], 2.5),
        ("B-1/1", "Фізика", ["Петров П. П."], 2),
        ("B-1/1", "Історія", ["Сидоренко С. С."], 1.5),
        ("B-1/1", "Іноземна мова", ["Коваль К. К."], 2),
        ("B-1/1", "Фізична культура", ["Бондар Б. Б."], 2),
        ("B-1/1", "Інформатика", ["Шевченко Т. Г."], 2),
        ("B-1/1", "Хімія", ["Мороз М. М."], 2),
        ("B-1/1", "Біологія", ["Коваль К. К."], 2),
    ]
    for group, subject, teachers, per_week in plan:
        lessons.append(Lesson(len(lessons), group, subject, teachers, per_week, 1))
    rooms = [Room(str(i)) for i in range(1, 6)] + [
        Room("21", kind="computer"),
        Room("спортзал", kind="gym", shared=True),
    ]
    return ProblemData(lessons, [Group("A-1/1", 1), Group("B-1/1", 1)], rooms)


def test_normalize_teacher():
    assert normalize_teacher("Сальчук А.В.") == "Сальчук А. В."
    assert normalize_teacher("  Качан  О.В. ") == "Качан О. В."


def test_small_problem_has_no_violations():
    data = small_problem()
    schedule = solve(data, FAST)
    assert schedule.status in ("OPTIMAL", "FEASIBLE")
    assign_rooms(schedule)
    issues = validate(schedule)
    assert [i for i in issues if i.severity == "error"] == []
    assert [i for i in issues if i.severity != "info"] == []

    counts = Counter((p.lesson_id, p.week) for p in schedule.placements)
    # Дробове навантаження чергується: в одному тижні менше, в іншому більше.
    for l in data.lessons:
        assert sorted((counts[l.id, 0], counts[l.id, 1])) == list(l.week_bounds)

    # Підгрупи отримують дві різні аудиторії.
    split = next(l for l in data.lessons if len(l.teachers) == 2)
    for p in schedule.placements:
        if p.lesson_id == split.id:
            assert len(set(p.rooms)) == 2


def test_second_shift_uses_pairs_5_to_7():
    data = small_problem()
    for g in data.groups:
        g.shift = 2
    for l in data.lessons:
        l.shift = 2
    schedule = solve(data, FAST)
    assert schedule.placements
    assert max(p.pair for p in schedule.placements) <= 7
    assert min(p.pair for p in schedule.placements) >= 4  # 4-та — лише як перехід між змінами


def test_allowed_pairs():
    assert Group("A", 1).allowed_pairs() == [1, 2, 3, 4, 5]
    assert Group("A", 2).allowed_pairs() == [4, 5, 6, 7]
    assert Group("A", 2).allowed_pairs(cross_shift=False) == [5, 6, 7]


def test_semester_matches_college_files():
    """Дати з файлів info/: 23.09 — знаменник, 26.09 — субота за четвергом, 28.09 — чисельник."""
    sem = Semester(date(2026, 9, 1), date(2026, 12, 31))
    assert sem.resolve(date(2026, 9, 23)) == (2, 1)
    assert sem.resolve(date(2026, 9, 26)) == (3, 1)
    assert sem.resolve(date(2026, 9, 28)) == (0, 0)
    assert sem.resolve(date(2026, 9, 27)) is None  # неділя
    sats = sem.saturdays_list()
    assert [s["day"] for s in sats[:7]] == [0, 1, 2, 3, 4, 0, 1]
    assert Semester(sem.start, sem.end, saturdays=False).resolve(date(2026, 9, 26)) is None


def test_day_workbook():
    schedule = solve(small_problem(), FAST)
    assign_rooms(schedule)
    schedule.semester = Semester(date(2026, 9, 1), date(2026, 12, 31))
    ws = build_day_workbook(schedule, date(2026, 9, 26)).active
    assert "четвер" in ws.cell(1, 1).value
    thursday = {p.pair for p in schedule.placements if p.day == 3 and p.week == 1 and p.lesson_id < 8}
    written = {ws.cell(3 + pair, 1).value for pair in range(1, 8) if ws.cell(3 + pair, 2).value}
    assert written == thursday


def test_unify_full_name_teacher():
    lessons = [
        Lesson(0, "A", "Захист України", ["Сергій ПАНДРАК"], 1, 1),
        Lesson(1, "B", "Захист України", ["Іван ПЕТРЕНКО"], 1, 1),
    ]
    warnings = []
    unify_teachers(lessons, ["Пандрак С. Б.", "Пандрак О. В."], warnings)
    assert lessons[0].teachers == ["Пандрак С. Б."]
    assert lessons[1].teachers == ["Петренко І."]
    assert len(warnings) == 2


def test_validator_detects_teacher_overlap():
    data = small_problem()
    schedule = solve(data, FAST)
    # Ставимо двом групам одного викладача в один час.
    a = next(p for p in schedule.placements if schedule.lesson(p.lesson_id).subject == "Математика")
    b_lesson = next(l for l in data.lessons if l.group == "B-1/1" and l.subject == "Математика")
    b = next(p for p in schedule.placements if p.lesson_id == b_lesson.id and p.week == a.week)
    b.day, b.pair = a.day, a.pair
    rules = {i.rule for i in validate(schedule) if i.severity == "error"}
    assert "Накладання викладача" in rules


def test_export_merges_identical_weeks():
    schedule = solve(small_problem(), FAST)
    assign_rooms(schedule)
    wb = build_workbook(schedule)
    assert wb.sheetnames[0] == "Розклад"
    assert "Викладачі" in wb.sheetnames and "Перевірка" in wb.sheetnames
    merged = [str(r) for r in wb["Розклад"].merged_cells.ranges]
    assert any(r.startswith("C") and ":C" in r for r in merged)


@pytest.mark.skipif(not SAMPLE.exists(), reason="немає прикладу вхідних даних")
def test_real_sample(tmp_path):
    schedule = generate(SAMPLE, SolverConfig(time_limit=30))
    assert schedule.placements
    errors = [i for i in schedule.issues if i.severity == "error"]
    assert errors == []
    out = tmp_path / "s.xlsx"
    build_workbook(schedule).save(out)
    assert openpyxl.load_workbook(out).sheetnames[:2] == ["Розклад", "1 курс"]


def test_impossible_teacher_load_is_explained(tmp_path):
    data = small_problem()
    # Один викладач на все: 30+ пар на тиждень — більше за 4 × 5.
    for l in data.lessons:
        l.teachers = ["Іванов І. І."]
    schedule = solve(data, SolverConfig(time_limit=5))
    assert not schedule.placements
    from app.scheduler.validator import precheck

    rules = {i.rule for i in precheck(data, n_days=5)}
    assert "Навантаження викладача" in rules
